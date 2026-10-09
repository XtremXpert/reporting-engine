# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html).

import io
import json
from urllib.parse import quote, urlencode

from openpyxl import load_workbook

from odoo.tests import HttpCase, tagged


@tagged("post_install", "-at_install")
class TestReportController(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.report = cls.env["ir.actions.report"].create(
            {
                "name": "Partner XLSX",
                "model": "res.partner",
                "report_type": "xlsx",
                "report_name": "report_xlsx.partner_xlsx",
                "print_report_name": "'Partner - %s' % object.name",
            }
        )
        cls.partner = cls.env["res.partner"].create({"name": "XLSX Partner"})

    def _download(self, url):
        query = urlencode(
            {"data": json.dumps([url, "xlsx"]), "context": json.dumps({})}
        )
        return self.url_open(f"/report/download?{query}")

    def _first_cell(self, response):
        workbook = load_workbook(io.BytesIO(response.content))
        return workbook.active.cell(1, 1).value

    def test_download_with_docids(self):
        """Generic report: records in the URL, named after print_report_name."""
        self.authenticate("admin", "admin")
        response = self._download(
            f"/report/xlsx/report_xlsx.partner_xlsx/{self.partner.id}?context={{}}"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            "Partner%20-%20XLSX%20Partner.xlsx",
            response.headers["Content-Disposition"],
        )
        self.assertEqual(self._first_cell(response), "XLSX Partner")

    def test_download_with_options(self):
        """Wizard-style report: records taken from the context in the URL."""
        self.authenticate("admin", "admin")
        options = quote(json.dumps({"foo": "bar"}))
        context = quote(json.dumps({"active_ids": [self.partner.id]}))
        response = self._download(
            f"/report/xlsx/report_xlsx.partner_xlsx?options={options}"
            f"&context={context}"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Partner%20XLSX.xlsx", response.headers["Content-Disposition"])
        self.assertEqual(self._first_cell(response), "XLSX Partner")
